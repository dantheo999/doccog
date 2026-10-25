from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.documents import DocumentError, DocumentStore, parse_document


def test_preserves_unicode_and_whitespace():
    text = "A \U0001f600 caf\u00e9\r\n  Alice signed.\n"
    assert parse_document("note.txt", text.encode()).text == text


def test_accepts_utf16_bom():
    assert parse_document("note.md", "Alice signed.".encode("utf-16")).text == "Alice signed."


@pytest.mark.parametrize("name,content", [
    ("notes.exe", b"text"), ("empty.txt", b" \n"), ("bad.txt", b"\xff\xfe\xff"),
    ("bad.docx", b"not a zip"), ("bad.pdf", b"not a pdf"),
])
def test_rejects_invalid_documents(name, content):
    with pytest.raises(DocumentError):
        parse_document(name, content)


def test_docx_preserves_paragraph_table_order():
    doc = Document()
    doc.add_paragraph("Before")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "In table"
    doc.add_paragraph("After")
    buffer = BytesIO()
    doc.save(buffer)
    assert parse_document("note.docx", buffer.getvalue()).text == "Before\nIn table\t\nAfter"


def test_rejects_pdf_without_text():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(DocumentError, match="OCR"):
        parse_document("scan.pdf", buffer.getvalue())


def test_pdf_extracts_selectable_text():
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=400)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 30 350 Td (Alice signed.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)
    parsed = parse_document("selectable.pdf", buffer.getvalue())
    assert parsed.text == "Alice signed."
    assert parsed.pages == 1


def test_rejects_encrypted_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("password")
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(DocumentError, match="Encrypted"):
        parse_document("protected.pdf", buffer.getvalue())


def test_rejects_upload_over_byte_limit():
    with pytest.raises(DocumentError) as error:
        parse_document("large.txt", b"a" * (20 * 1024 * 1024 + 1))
    assert error.value.status == 413


def test_rejects_text_over_limit():
    with pytest.raises(DocumentError, match="character"):
        parse_document("large.txt", b"a" * 500_001)


def test_store_expires_documents():
    clock = [0.0]
    store = DocumentStore(ttl=10, clock=lambda: clock[0])
    document = store.add(parse_document("note.txt", b"Alice"))
    assert store.get(document.id).text == "Alice"
    clock[0] = 11
    with pytest.raises(DocumentError, match="expired"):
        store.get(document.id)


def test_store_evicts_old_documents_at_capacity():
    store = DocumentStore(max_documents=1)
    first = store.add(parse_document("one.txt", b"Alice"))
    second = store.add(parse_document("two.txt", b"Bob"))
    assert store.get(second.id).text == "Bob"
    with pytest.raises(DocumentError):
        store.get(first.id)
